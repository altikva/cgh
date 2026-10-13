# -#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#
# __creation__ = 2026-10-13
# __author__ = "jndjama (Joy Ndjama)"
# __copyright__ = "Copyright 2026 ALTIKVA."
# __licence__ = "Apache-2.0 (https://www.apache.org/licenses/LICENSE-2.0)"
# -#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#
# Description: Document parsers for pdf, docx, xlsx and pptx. Each file
#              becomes searchable sections, exactly like Markdown headings,
#              so search_docs, doc_outline, fts_search and the federated
#              fan-out work on them unchanged. The extracted text is also
#              exposed as idx.scan_text so scanners see page / cell / slide
#              content rather than the raw zip or pdf bytes.
#
#              Parsing is best effort: an encrypted pdf or a corrupt workbook
#              yields an empty index and one log line, never a failed scan.
#              The reading libraries are imported lazily inside parse(), so
#              a build that lacks one of them (a trimmed binary, a broken
#              lxml wheel) keeps every parser registered and skips that format
#              with a single log line per process instead of crashing.
#
#              Big files stay cheap: the indexer's max_file_size_kb cap and
#              ignore_patterns apply before any of this runs, and each parser
#              bounds its own work (pages, rows, slides, extracted characters)
#              so a large file that is allowed through costs a bounded parse.

from __future__ import annotations

import logging
import re
import warnings
from pathlib import Path

from . import register_parser
from .base import BaseParser, FileIndex, SectionDef

_log = logging.getLogger(__name__)

_PREVIEW_CHARS = 400  # what lands in the FTS per section
_MAX_SCAN_CHARS = 200_000  # extracted text handed to scanners, per file
_MAX_PDF_PAGES = 500  # pages turned into sections; the rest is skipped
_MAX_XLSX_ROWS = 50_000  # rows walked per workbook, empty rows included
_MAX_HEADER_CELLS = 30
_MAX_SLIDES = 500
_HEADING_RE = re.compile(r"heading\s*(\d)", re.IGNORECASE)

_missing_logged: set[str] = set()


def _missing_dependency(fmt: str, module: str, exc: BaseException) -> None:
    """Log once per process that a format's reader is not importable."""
    if fmt in _missing_logged:
        return
    _missing_logged.add(fmt)
    _log.warning(
        "%s files skipped: %s is not importable in this install (%s)",
        fmt,
        module,
        exc,
    )


class _ScanText:
    """Bounded accumulator for the scanner text of one file."""

    def __init__(self, limit: int = _MAX_SCAN_CHARS) -> None:
        self.parts: list[str] = []
        self.size = 0
        self.limit = limit

    @property
    def full(self) -> bool:
        return self.size >= self.limit

    def add(self, text: str) -> None:
        if text and not self.full:
            self.parts.append(text)
            self.size += len(text) + 1

    def value(self) -> str:
        return "\n".join(self.parts)[: self.limit]


# ---------------------------------------------------------------------------
# pdf
# ---------------------------------------------------------------------------


@register_parser(".pdf")
class PdfParser(BaseParser):
    """PDF pages and outline entries as document sections."""

    lang = "pdf"
    extensions = [".pdf"]
    extracts = ["sections"]
    description = "PDF pages and outline entries as document sections"

    def parse(self, path: Path) -> FileIndex:
        # One section per page with the page text as preview. Top-level
        # outline titles name the page they point at. "Lines" are page
        # numbers: Read cannot open a pdf at a line, and a page is what a
        # human asks for.
        idx = FileIndex(path=str(path), lang=self.lang)
        try:
            from pypdf import PdfReader
        except Exception as exc:
            _missing_dependency("pdf", "pypdf", exc)
            return idx
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    return idx

            outline_titles = _pdf_outline_by_page(reader)
            scan = _ScanText()
            total = len(reader.pages)
            for page_no, page in enumerate(reader.pages, start=1):
                if page_no > _MAX_PDF_PAGES:
                    _log.info(
                        "pdf %s: %d pages, only the first %d indexed",
                        path,
                        total,
                        _MAX_PDF_PAGES,
                    )
                    break
                try:
                    text = (page.extract_text() or "").strip()
                except Exception:
                    text = ""
                scan.add(text)
                title = outline_titles.get(page_no) or f"Page {page_no}"
                idx.sections.append(
                    SectionDef(
                        id=f"{path}::page{page_no}",
                        title=title,
                        level=1 if page_no in outline_titles else 2,
                        file_path=str(path),
                        start_line=page_no,
                        end_line=page_no,
                        body_preview=text[:_PREVIEW_CHARS],
                        anchor=f"page-{page_no}",
                    )
                )
            idx.scan_text = scan.value()
        except Exception:
            _log.warning("pdf parse failed, indexed empty: %s", path)
            return FileIndex(path=str(path), lang=self.lang)
        return idx


def _pdf_outline_by_page(reader) -> dict[int, str]:
    """Map page number to the first top-level outline title on it."""
    titles: dict[int, str] = {}
    try:
        for item in reader.outline:
            if isinstance(item, list):
                continue  # nested outline levels, keep it flat
            try:
                page_no = reader.get_destination_page_number(item) + 1
                titles.setdefault(page_no, str(item.title))
            except Exception:
                continue
    except Exception:
        pass
    return titles


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------


@register_parser(".docx")
class DocxParser(BaseParser):
    """Word headings and body text as document sections."""

    lang = "docx"
    extensions = [".docx"]
    extracts = ["sections"]
    description = "Word headings and body text as document sections"

    def parse(self, path: Path) -> FileIndex:
        # Heading-styled paragraphs become the section tree (mirroring the
        # markdown parser); the body text between headings is each
        # section's preview. "Lines" are paragraph indexes, the closest
        # stable anchor a docx has.
        idx = FileIndex(path=str(path), lang=self.lang)
        try:
            import docx
        except Exception as exc:
            _missing_dependency("docx", "python-docx", exc)
            return idx
        try:
            document = docx.Document(str(path))
        except Exception:
            _log.warning("docx parse failed, indexed empty: %s", path)
            return idx

        paragraphs = document.paragraphs
        current: SectionDef | None = None
        body: list[str] = []
        preamble: list[str] = []
        scan = _ScanText()

        def _flush() -> None:
            if current is not None:
                current.body_preview = " ".join(body)[:_PREVIEW_CHARS]

        for i, para in enumerate(paragraphs, start=1):
            text = (para.text or "").strip()
            scan.add(text)
            style = getattr(getattr(para, "style", None), "name", "") or ""
            m = _HEADING_RE.search(style)
            if m and text:
                _flush()
                if current is not None:
                    current.end_line = i - 1
                current = SectionDef(
                    id=f"{path}::para{i}",
                    title=text,
                    level=int(m.group(1)),
                    file_path=str(path),
                    start_line=i,
                    end_line=i,
                    anchor=f"para-{i}",
                )
                idx.sections.append(current)
                body = []
            elif text:
                (body if current is not None else preamble).append(text)

        _flush()
        if current is not None:
            current.end_line = len(paragraphs)

        # Table cells too: data often sits in tables, which the paragraph
        # walk above never visits.
        table_text: list[str] = []
        try:
            for table in document.tables:
                for trow in table.rows:
                    cells = [
                        c.text.strip() for c in trow.cells if (c.text or "").strip()
                    ]
                    if cells:
                        line = "\t".join(cells)
                        scan.add(line)
                        table_text.append(" | ".join(cells))
        except Exception:
            pass
        idx.scan_text = scan.value()

        # A document with no heading styles still gets one section so its
        # text (or, failing any prose, its tables) is searchable.
        if not idx.sections and (preamble or table_text):
            idx.sections.append(
                SectionDef(
                    id=f"{path}::body",
                    title=path.stem,
                    level=1,
                    file_path=str(path),
                    start_line=1,
                    end_line=max(len(paragraphs), 1),
                    body_preview=" ".join(preamble or table_text)[:_PREVIEW_CHARS],
                    anchor="body",
                )
            )
        return idx


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------


@register_parser(".xlsx", ".xlsm")
class XlsxParser(BaseParser):
    """Excel sheets with their header row as document sections."""

    lang = "xlsx"
    extensions = [".xlsx", ".xlsm"]
    extracts = ["sections"]
    description = "Excel sheets with their header row as document sections"

    def parse(self, path: Path) -> FileIndex:
        # One section per sheet: header row and sheet size in the preview
        # so column names are searchable. Read-only, values only, and a
        # row budget across the workbook, so a big file costs little.
        idx = FileIndex(path=str(path), lang=self.lang)
        try:
            import openpyxl
        except Exception as exc:
            _missing_dependency("xlsx", "openpyxl", exc)
            return idx
        wb = None
        try:
            with warnings.catch_warnings():
                # "Workbook contains no default style": cosmetic, and it
                # lands on the owner's stderr for every such file.
                warnings.simplefilter("ignore", UserWarning)
                wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            scan = _ScanText()
            rows_left = _MAX_XLSX_ROWS
            for sheet_no, ws in enumerate(wb.worksheets, start=1):
                header: list[str] = []
                if rows_left > 0 and not scan.full:
                    scan.add(f"# {ws.title}")
                    try:
                        for row_no, row in enumerate(ws.iter_rows(values_only=True)):
                            rows_left -= 1
                            cells = [str(c) for c in row if c is not None]
                            if row_no == 0:
                                header = [
                                    str(c)
                                    for c in row[:_MAX_HEADER_CELLS]
                                    if c is not None
                                ]
                            if cells:
                                scan.add("\t".join(cells))
                            if rows_left <= 0 or scan.full:
                                break
                    except Exception:
                        pass
                preview = f"{ws.max_row or 0} rows x {ws.max_column or 0} cols" + (
                    f" | columns: {', '.join(header)}" if header else ""
                )
                idx.sections.append(
                    SectionDef(
                        id=f"{path}::sheet{sheet_no}",
                        title=ws.title,
                        level=1,
                        file_path=str(path),
                        start_line=sheet_no,
                        end_line=sheet_no,
                        body_preview=preview[:_PREVIEW_CHARS],
                        anchor=f"sheet-{sheet_no}",
                    )
                )
            idx.scan_text = scan.value()
        except Exception:
            _log.warning("xlsx parse failed, indexed empty: %s", path)
            return FileIndex(path=str(path), lang=self.lang)
        finally:
            try:
                if wb is not None:
                    wb.close()
            except Exception:
                pass
        return idx


# ---------------------------------------------------------------------------
# pptx
# ---------------------------------------------------------------------------


@register_parser(".pptx")
class PptxParser(BaseParser):
    """PowerPoint slides, speaker notes and tables as document sections."""

    lang = "pptx"
    extensions = [".pptx"]
    extracts = ["sections"]
    description = "PowerPoint slides, speaker notes and tables as document sections"

    def parse(self, path: Path) -> FileIndex:
        # One level-1 section per slide, titled by the slide's title
        # placeholder (else "Slide N"), with its text frames and table
        # rows as the preview. Speaker notes get a level-2 section of their
        # own so they stay searchable even when the slide text is long.
        # "Lines" are slide numbers.
        idx = FileIndex(path=str(path), lang=self.lang)
        try:
            from pptx import Presentation
        except Exception as exc:
            _missing_dependency("pptx", "python-pptx", exc)
            return idx
        try:
            prs = Presentation(str(path))
            scan = _ScanText()
            slides = list(prs.slides)
            for slide_no, slide in enumerate(slides, start=1):
                if slide_no > _MAX_SLIDES:
                    _log.info(
                        "pptx %s: %d slides, only the first %d indexed",
                        path,
                        len(slides),
                        _MAX_SLIDES,
                    )
                    break
                title, lines = _pptx_slide_text(slide)
                notes = _pptx_notes(slide)
                body = " ".join(lines)
                scan.add(f"# Slide {slide_no}")
                for line in lines:
                    scan.add(line)
                idx.sections.append(
                    SectionDef(
                        id=f"{path}::slide{slide_no}",
                        title=title or f"Slide {slide_no}",
                        level=1,
                        file_path=str(path),
                        start_line=slide_no,
                        end_line=slide_no,
                        body_preview=body[:_PREVIEW_CHARS],
                        anchor=f"slide-{slide_no}",
                    )
                )
                if notes:
                    scan.add(notes)
                    idx.sections.append(
                        SectionDef(
                            id=f"{path}::slide{slide_no}-notes",
                            title=f"Notes: {title or f'Slide {slide_no}'}",
                            level=2,
                            file_path=str(path),
                            start_line=slide_no,
                            end_line=slide_no,
                            body_preview=notes[:_PREVIEW_CHARS],
                            anchor=f"slide-{slide_no}-notes",
                        )
                    )
            idx.scan_text = scan.value()
        except Exception:
            _log.warning("pptx parse failed, indexed empty: %s", path)
            return FileIndex(path=str(path), lang=self.lang)
        return idx


def _pptx_slide_text(slide) -> tuple[str, list[str]]:
    """The slide's title and its other text, one entry per paragraph or
    table row, in shape order. Group shapes are walked."""
    title = ""
    try:
        if slide.shapes.title is not None:
            title = (slide.shapes.title.text_frame.text or "").strip()
    except Exception:
        title = ""
    lines: list[str] = []

    def _walk(shapes) -> None:
        for shape in shapes:
            try:
                if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
                    _walk(shape.shapes)
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        cells = [
                            (c.text or "").strip()
                            for c in row.cells
                            if (c.text or "").strip()
                        ]
                        if cells:
                            lines.append(" | ".join(cells))
                    continue
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = "".join(r.text for r in para.runs).strip()
                        if text and text != title:
                            lines.append(text)
            except Exception:
                continue

    _walk(slide.shapes)
    return title, lines


def _pptx_notes(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
        return " ".join((frame.text or "").split()) if frame is not None else ""
    except Exception:
        return ""
