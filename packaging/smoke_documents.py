# -#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#
# __creation__ = 2026-10-13
# __author__ = "jndjama (Joy Ndjama)"
# __copyright__ = "Copyright 2026 ALTIKVA."
# __licence__ = "Apache-2.0 (https://www.apache.org/licenses/LICENSE-2.0)"
# -#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#
# Description: Document smoke test for a built cgh (the standalone binary or
#              any `cgh` on PATH): writes one pdf, docx, xlsx, pptx and csv
#              into a scratch git repo, runs `cgh index` there, then requires
#              `cgh search --text` to find a word only that file holds. A
#              binary missing a document reader still starts (the parsers
#              skip a format whose reader is absent), so this is what proves
#              the readers were bundled.
#
#              Usage: python smoke_documents.py <cgh executable> [workdir]
#              The fixture writers need python-docx, openpyxl and
#              python-pptx in the Python running this script, e.g.
#                uv run --no-project --with python-docx --with openpyxl \
#                  --with python-pptx python packaging/smoke_documents.py cgh

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

# word -> file that holds it, and nothing else does
_EXPECT = {
    "pelican": "invoice.pdf",
    "marmoset": "contract.docx",
    "aardvark": "budget.xlsx",
    "walrus": "deck.pptx",
    "narwhal": "donors.csv",
}


def _pdf(text: str) -> bytes:
    stream = b"BT /F1 24 Tf 72 700 Td (%s) Tj ET" % text.encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (i, body)
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Root 1 0 R /Size %d >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_at,
    )
    return bytes(out)


def write_fixtures(root: Path) -> None:
    import docx
    import openpyxl
    from pptx import Presentation

    (root / "invoice.pdf").write_bytes(_pdf("Invoice for the pelican sanctuary"))

    d = docx.Document()
    d.add_heading("Supply contract", level=1)
    d.add_paragraph("The marmoset enclosure is delivered within thirty days.")
    d.save(str(root / "contract.docx"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["aardvark_feed", "amount"])
    ws.append(["hay", 120])
    wb.save(str(root / "budget.xlsx"))

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Open the new wing"
    slide.notes_slide.notes_text_frame.text = "Mention the walrus budget"
    prs.save(str(root / "deck.pptx"))

    (root / "donors.csv").write_text(
        "donor,city\nnarwhal foundation,Brest\nAda,Lyon\n", encoding="utf-8"
    )


def main(argv: list[str]) -> int:
    cgh = argv[1] if len(argv) > 1 else "cgh"
    root = Path(argv[2]) if len(argv) > 2 else Path(tempfile.mkdtemp(prefix="cgh-"))
    root.mkdir(parents=True, exist_ok=True)
    subprocess.run(["git", "init", "-q"], cwd=root, check=True)
    write_fixtures(root)

    run = subprocess.run(
        [cgh, "index", "--root", str(root)],
        cwd=root,
        capture_output=True,
        text=True,
        timeout=300,
    )
    print(run.stdout[-2000:])
    if run.returncode != 0:
        print(run.stderr[-2000:], file=sys.stderr)
        return 1

    failed = []
    for word, name in _EXPECT.items():
        out = subprocess.run(
            [cgh, "search", "--text", word, "--root", str(root)],
            cwd=root,
            capture_output=True,
            text=True,
            timeout=120,
        )
        found = name in out.stdout
        print(f"{'ok  ' if found else 'FAIL'} {word!r} -> {name}")
        if not found:
            print(out.stdout[-1000:], out.stderr[-1000:], file=sys.stderr)
            failed.append(name)
    if failed:
        print(f"documents not searchable: {failed}", file=sys.stderr)
        return 1
    print("every document format is indexed and searchable")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
