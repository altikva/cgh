# -#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#
# __creation__ = 2026-10-13
# __author__ = "jndjama (Joy Ndjama)"
# __copyright__ = "Copyright 2026 ALTIKVA."
# __licence__ = "Apache-2.0 (https://www.apache.org/licenses/LICENSE-2.0)"
# -#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#-#
# Description: Core document parsers: pdf pages, docx heading tree and
#              tables, xlsx sheets, pptx slides / notes / tables, csv
#              summaries. Covers corrupt-file resilience, the work bounds,
#              the missing-library fallback, scanner text, the end-to-end
#              chain (index a repo, find the sections through the FTS) and
#              that an incremental index picks up document files an older
#              cgh skipped, with no graph format bump. Every fixture is
#              generated in the test.

from __future__ import annotations

import logging
import subprocess
import sys

import pytest

import codegraph.parsers as parsers
import codegraph.parsers.csv_table as csv_table
import codegraph.parsers.documents as documents
from codegraph.parsers import get_parser_for_path
from codegraph.parsers.csv_table import CsvParser
from codegraph.parsers.documents import DocxParser, PdfParser, PptxParser, XlsxParser

# ---------------------------------------------------------------------------
# fixtures, built on the fly
# ---------------------------------------------------------------------------


def _mini_pdf(text: str = "Hello cgh docs") -> bytes:
    """A complete single-page PDF with one Helvetica text object and a
    correct xref table, so pypdf accepts it. Handwritten because pypdf
    writes but does not typeset."""
    stream = b"BT /F1 24 Tf 72 700 Td (%s) Tj ET" % text.encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (i, body)
    xref_at = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Root 1 0 R /Size %d >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_at,
    )
    return bytes(out)


def _make_docx(path):
    import docx

    d = docx.Document()
    d.add_heading("Contract terms", level=1)
    d.add_paragraph("The supplier delivers within thirty days.")
    d.add_heading("Payment", level=2)
    d.add_paragraph("Net sixty, in euros.")
    d.save(str(path))


def _make_xlsx(path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Donations"
    ws.append(["donor_email", "amount", "currency"])
    ws.append(["a@b.c", 100, "EUR"])
    wb.create_sheet("Refunds").append(["ref_id", "reason"])
    wb.save(str(path))


def _make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])  # title and content
    s1.shapes.title.text = "Quarterly roadmap"
    s1.placeholders[1].text = "Ship the zebra importer"
    s1.notes_slide.notes_text_frame.text = "Mention the walrus budget"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s2.shapes.title.text = "Headcount"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "team"
    table.cell(0, 1).text = "hires"
    table.cell(1, 0).text = "platypus squad"
    table.cell(1, 1).text = "7"
    prs.slides.add_slide(prs.slide_layouts[6])  # blank, untitled
    prs.save(str(path))


# ---------------------------------------------------------------------------
# dispatch
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "cls"),
    [
        ("a.pdf", "PdfParser"),
        ("a.docx", "DocxParser"),
        ("a.xlsx", "XlsxParser"),
        ("a.xlsm", "XlsxParser"),
        ("a.pptx", "PptxParser"),
        ("a.csv", "CsvParser"),
        ("a.tsv", "CsvParser"),
    ],
)
def test_core_claims_document_extensions(name, cls, tmp_path):
    parser = get_parser_for_path(tmp_path / name)
    assert parser is not None and type(parser).__name__ == cls
    assert parsers.is_supported(tmp_path / name)


# ---------------------------------------------------------------------------
# pdf / docx / xlsx (ported from the cgh-docs plugin)
# ---------------------------------------------------------------------------


class TestPdfParser:
    def test_extracts_page_sections(self, tmp_path):
        f = tmp_path / "doc.pdf"
        f.write_bytes(_mini_pdf())
        idx = PdfParser().parse(f)
        assert idx.lang == "pdf"
        assert len(idx.sections) == 1
        assert idx.sections[0].title == "Page 1"
        assert "Hello cgh docs" in idx.sections[0].body_preview
        assert "Hello cgh docs" in idx.scan_text

    def test_corrupt_pdf_yields_empty_index(self, tmp_path):
        f = tmp_path / "bad.pdf"
        f.write_bytes(b"%PDF-not really")
        assert PdfParser().parse(f).sections == []

    def test_page_cap(self, tmp_path, monkeypatch):
        from pypdf import PdfReader, PdfWriter

        one = tmp_path / "one.pdf"
        one.write_bytes(_mini_pdf())
        writer = PdfWriter()
        for _ in range(5):
            writer.add_page(PdfReader(str(one)).pages[0])
        many = tmp_path / "many.pdf"
        with open(many, "wb") as fh:
            writer.write(fh)
        monkeypatch.setattr(documents, "_MAX_PDF_PAGES", 2)
        assert len(PdfParser().parse(many).sections) == 2

        # Pages past the cap are never extracted, not extracted then dropped.
        from pypdf import PageObject

        calls = []
        real = PageObject.extract_text

        def spy(self, *a, **k):
            calls.append(1)
            return real(self, *a, **k)

        monkeypatch.setattr(PageObject, "extract_text", spy)
        PdfParser().parse(many)
        assert len(calls) == 2

    def test_text_caps_per_page_and_per_file(self, tmp_path, monkeypatch):
        from pypdf import PdfReader, PdfWriter

        one = tmp_path / "one.pdf"
        one.write_bytes(_mini_pdf("Hello cgh docs"))
        writer = PdfWriter()
        for _ in range(4):
            writer.add_page(PdfReader(str(one)).pages[0])
        many = tmp_path / "many.pdf"
        with open(many, "wb") as fh:
            writer.write(fh)

        monkeypatch.setattr(documents, "_MAX_PAGE_CHARS", 5)
        monkeypatch.setattr(documents, "_MAX_PDF_TEXT_CHARS", 10)
        idx = PdfParser().parse(many)
        previews = [s.body_preview for s in idx.sections]
        assert previews == ["Hello", "Hello", "", ""]


class TestDocxParser:
    def test_heading_tree_and_previews(self, tmp_path):
        f = tmp_path / "contract.docx"
        _make_docx(f)
        idx = DocxParser().parse(f)
        assert [(s.title, s.level) for s in idx.sections] == [
            ("Contract terms", 1),
            ("Payment", 2),
        ]
        assert "thirty days" in idx.sections[0].body_preview
        assert "euros" in idx.sections[1].body_preview

    def test_table_cells_reach_scan_text(self, tmp_path):
        import docx

        document = docx.Document()
        document.add_paragraph("Contact list")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Bob Durand"
        table.rows[0].cells[1].text = "+33 6 11 22 33 44"
        path = tmp_path / "contacts.docx"
        document.save(str(path))

        idx = DocxParser().parse(path)
        assert "Bob Durand" in idx.scan_text
        assert "+33 6 11 22 33 44" in idx.scan_text
        # No headings: one body section keeps the prose searchable.
        assert [s.title for s in idx.sections] == ["contacts"]

    def test_corrupt_docx_yields_empty_index(self, tmp_path):
        f = tmp_path / "bad.docx"
        f.write_bytes(b"not a zip at all")
        assert DocxParser().parse(f).sections == []


class TestXlsxParser:
    def test_sheets_with_headers(self, tmp_path):
        f = tmp_path / "book.xlsx"
        _make_xlsx(f)
        idx = XlsxParser().parse(f)
        assert [s.title for s in idx.sections] == ["Donations", "Refunds"]
        assert "donor_email" in idx.sections[0].body_preview
        assert "ref_id" in idx.sections[1].body_preview

    def test_data_cells_reach_scan_text(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name", "email"])
        ws.append(["Alice Martin", "alice.martin@example.com"])
        path = tmp_path / "people.xlsx"
        wb.save(str(path))

        idx = XlsxParser().parse(path)
        assert "alice.martin@example.com" in idx.scan_text
        assert "Alice Martin" in idx.scan_text

    def test_row_budget_bounds_the_walk(self, tmp_path, monkeypatch):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["id", "value"])
        for i in range(200):
            ws.append([i, f"row-{i}"])
        path = tmp_path / "big.xlsx"
        wb.save(str(path))

        idx = XlsxParser().parse(path)
        assert idx.sections[0].body_preview.startswith("201 rows x 2 cols")

        monkeypatch.setattr(documents, "_MAX_XLSX_ROWS", 10)
        idx = XlsxParser().parse(path)
        assert "row-5" in idx.scan_text
        assert "row-150" not in idx.scan_text
        # The size shown is what was walked, marked as a lower bound.
        assert idx.sections[0].body_preview.startswith("10+ rows x 2 cols")

    def test_declared_dimension_is_not_trusted(self, tmp_path, monkeypatch):
        # A sheet declaring A1:XFD1048576 must not drive the walk: only the
        # rows actually streamed count, and the cell budget stops wide rows.
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        for i in range(50):
            ws.append([f"c{i}-{j}" for j in range(40)])
        path = tmp_path / "wide.xlsx"
        wb.save(str(path))

        monkeypatch.setattr(documents, "_MAX_XLSX_CELLS", 400)
        idx = XlsxParser().parse(path)
        assert idx.sections[0].body_preview.startswith("10+ rows x 40 cols")
        assert "c9-0" in idx.scan_text and "c10-0" not in idx.scan_text

    def test_sheet_count_cap(self, tmp_path, monkeypatch):
        import openpyxl

        wb = openpyxl.Workbook()
        for i in range(6):
            wb.create_sheet(f"S{i}")
        path = tmp_path / "sheets.xlsx"
        wb.save(str(path))
        monkeypatch.setattr(documents, "_MAX_XLSX_SHEETS", 3)
        assert len(XlsxParser().parse(path).sections) == 3

    def test_corrupt_xlsx_yields_empty_index(self, tmp_path):
        f = tmp_path / "bad.xlsx"
        f.write_bytes(b"nope")
        assert XlsxParser().parse(f).sections == []


# ---------------------------------------------------------------------------
# pptx
# ---------------------------------------------------------------------------


class TestPptxParser:
    def test_slides_notes_and_tables(self, tmp_path):
        f = tmp_path / "deck.pptx"
        _make_pptx(f)
        idx = PptxParser().parse(f)
        assert idx.lang == "pptx"
        assert [(s.title, s.level, s.start_line) for s in idx.sections] == [
            ("Quarterly roadmap", 1, 1),
            ("Notes: Quarterly roadmap", 2, 1),
            ("Headcount", 1, 2),
            ("Slide 3", 1, 3),
        ]
        by_title = {s.title: s for s in idx.sections}
        assert "zebra importer" in by_title["Quarterly roadmap"].body_preview
        assert "walrus budget" in by_title["Notes: Quarterly roadmap"].body_preview
        assert "platypus squad | 7" in by_title["Headcount"].body_preview
        for needle in ("zebra importer", "walrus budget", "platypus squad"):
            assert needle in idx.scan_text

    def test_slide_cap(self, tmp_path, monkeypatch):
        f = tmp_path / "deck.pptx"
        _make_pptx(f)
        monkeypatch.setattr(documents, "_MAX_SLIDES", 1)
        idx = PptxParser().parse(f)
        assert [s.start_line for s in idx.sections] == [1, 1]

    def test_slides_past_the_cap_are_never_read(self, tmp_path, monkeypatch):
        from pptx import Presentation

        prs = Presentation()
        for i in range(12):
            prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = f"T{i}"
        f = tmp_path / "long.pptx"
        prs.save(str(f))

        seen = []
        real = documents._pptx_slide_text

        def spy(slide, budget):
            seen.append(slide)
            return real(slide, budget)

        monkeypatch.setattr(documents, "_pptx_slide_text", spy)
        monkeypatch.setattr(documents, "_MAX_SLIDES", 4)
        idx = PptxParser().parse(f)
        assert [s.title for s in idx.sections] == ["T0", "T1", "T2", "T3"]
        assert len(seen) == 4

    def test_shape_budget_spans_the_deck(self, tmp_path, monkeypatch):
        f = tmp_path / "deck.pptx"
        _make_pptx(f)
        monkeypatch.setattr(documents, "_MAX_PPTX_SHAPES", 2)
        idx = PptxParser().parse(f)
        # Slide 1 spends the budget (title + body), so the table on slide 2
        # is never walked; its title still comes from the placeholder.
        assert "platypus" not in idx.scan_text
        assert "zebra importer" in idx.scan_text

    def test_corrupt_pptx_yields_empty_index(self, tmp_path):
        f = tmp_path / "bad.pptx"
        f.write_bytes(b"PK not really a deck")
        assert PptxParser().parse(f).sections == []


# ---------------------------------------------------------------------------
# csv
# ---------------------------------------------------------------------------


class TestCsvParser:
    def test_header_size_and_first_rows(self, tmp_path):
        f = tmp_path / "donors.csv"
        f.write_text(
            "donor,amount,city\nAda,10,Lyon\nGrace,20,Nantes\nLin,30,Lille\n"
            "Max,40,Brest\n",
            encoding="utf-8",
        )
        idx = CsvParser().parse(f)
        assert idx.lang == "csv"
        assert len(idx.sections) == 1
        sec = idx.sections[0]
        assert sec.title == "donors.csv"
        assert sec.body_preview.startswith("4 rows x 3 cols | columns: donor, amount")
        assert "Grace ; 20 ; Nantes" in sec.body_preview
        assert "Max" not in sec.body_preview  # only the first rows
        assert "Max\t40\tBrest" in idx.scan_text  # scanners see them all

    def test_semicolon_and_tsv(self, tmp_path):
        semi = tmp_path / "fr.csv"
        semi.write_text("nom;ville\nAda;Lyon\n", encoding="utf-8")
        assert "columns: nom, ville" in CsvParser().parse(semi).sections[0].body_preview
        tsv = tmp_path / "t.tsv"
        tsv.write_text("a\tb\n1\t2\n", encoding="utf-8")
        assert "columns: a, b" in CsvParser().parse(tsv).sections[0].body_preview

    def test_huge_csv_is_bounded(self, tmp_path, monkeypatch):
        f = tmp_path / "huge.csv"
        f.write_text(
            "id,payload\n" + "".join(f"{i},value-{i}\n" for i in range(5000)),
            encoding="utf-8",
        )
        monkeypatch.setattr(csv_table, "_MAX_ROWS", 100)
        idx = CsvParser().parse(f)
        assert idx.sections[0].body_preview.startswith("5000 rows x 2 cols")
        assert "value-50\n" in idx.scan_text
        assert "value-4000" not in idx.scan_text

        # Past the counting budget the row count turns into a lower bound.
        monkeypatch.setattr(csv_table, "_COUNT_BYTES", 1024)
        preview = CsvParser().parse(f).sections[0].body_preview
        assert preview.split(" ", 1)[0].endswith("+")

    def test_a_line_without_newline_is_never_read_whole(self, tmp_path, monkeypatch):
        f = tmp_path / "oneline.csv"
        f.write_text("a,b\n1,2\n" + "x," * 600_000, encoding="utf-8")
        monkeypatch.setattr(csv_table, "_READ_CHARS", 1024)
        reads = []
        real_open = open

        class Spy:
            def __init__(self, fh):
                self.fh = fh

            def __enter__(self):
                return self

            def __exit__(self, *exc):
                self.fh.close()

            def read(self, n=-1):
                assert n != -1, "unbounded read"
                reads.append(n)
                return self.fh.read(n)

        def fake_open(p, *a, **k):
            fh = real_open(p, *a, **k)
            return Spy(fh) if "b" not in (a[0] if a else k.get("mode", "r")) else fh

        monkeypatch.setattr("builtins.open", fake_open)
        idx = CsvParser().parse(f)
        assert reads and max(reads) <= 1024
        # The cut row is dropped; the two complete ones are kept.
        assert idx.scan_text == "a\tb\n1\t2"

    @pytest.mark.skipif(not hasattr(__import__("os"), "mkfifo"), reason="POSIX only")
    def test_not_a_regular_file(self, tmp_path):
        import os

        fifo = tmp_path / "pipe.csv"
        os.mkfifo(fifo)
        assert CsvParser().parse(fifo).sections == []  # would block if opened

    def test_empty_and_unreadable(self, tmp_path):
        empty = tmp_path / "empty.csv"
        empty.write_text("", encoding="utf-8")
        idx = CsvParser().parse(empty)
        assert idx.sections[0].body_preview.startswith("0 rows")
        assert CsvParser().parse(tmp_path).sections == []  # a directory


# ---------------------------------------------------------------------------
# zip containers: the directory is checked before any library opens them
# ---------------------------------------------------------------------------

_ZIP_FORMATS = [
    ("docx", "word/document.xml"),
    ("xlsx", "xl/worksheets/sheet1.xml"),
    ("pptx", "ppt/slides/slide1.xml"),
]


@pytest.fixture
def library_calls(monkeypatch):
    """Record every call into the three office libraries' openers."""
    import docx
    import openpyxl
    import pptx

    calls: list[str] = []

    def recorder(name):
        def _call(*a, **k):
            calls.append(name)
            raise AssertionError(f"{name} must not be reached")

        return _call

    monkeypatch.setattr(docx, "Document", recorder("docx"))
    monkeypatch.setattr(openpyxl, "load_workbook", recorder("xlsx"))
    monkeypatch.setattr(pptx, "Presentation", recorder("pptx"))
    return calls


def _parse(fmt, path):
    return {"docx": DocxParser, "xlsx": XlsxParser, "pptx": PptxParser}[fmt]().parse(
        path
    )


def _write_zip(path, members):
    """members: name -> bytes, or name -> int for that many zero bytes,
    streamed in 1 MB chunks so a bomb never sits in memory whole."""
    import zipfile

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            if isinstance(data, int):
                with zf.open(name, "w", force_zip64=True) as out:
                    chunk = b"\0" * (1 << 20)
                    for _ in range(data >> 20):
                        out.write(chunk)
            else:
                zf.writestr(name, data)


@pytest.fixture
def doc_log():
    """Warnings the document parsers log, collected by a handler attached
    to their own logger. caplog listens on the root logger, so it misses
    these whenever something has set propagate=False on "codegraph" (the
    background logging setup does, by design); this does not depend on
    propagation, nor on the level a parent logger was left at."""
    records: list[logging.LogRecord] = []

    class _Collect(logging.Handler):
        def emit(self, record):
            records.append(record)

    logger = logging.getLogger("codegraph.parsers.documents")
    handler = _Collect(logging.WARNING)
    saved = (logger.level, logger.disabled)
    logger.addHandler(handler)
    logger.setLevel(logging.WARNING)
    logger.disabled = False
    try:
        yield records
    finally:
        logger.removeHandler(handler)
        logger.setLevel(saved[0])
        logger.disabled = saved[1]


@pytest.mark.parametrize(("fmt", "member"), _ZIP_FORMATS)
def test_zip_bomb_is_skipped_before_the_library(
    tmp_path, library_calls, doc_log, fmt, member
):
    import time

    bomb = tmp_path / f"bomb.{fmt}"
    # 100 MB of zeros, about 100 KB on disk: well under max_file_size_kb.
    _write_zip(bomb, {"[Content_Types].xml": b"<Types/>", member: 100 << 20})
    assert bomb.stat().st_size < 500 * 1024

    t0 = time.perf_counter()
    idx = _parse(fmt, bomb)
    assert time.perf_counter() - t0 < 1.0
    assert idx.sections == [] and idx.scan_text == ""
    assert library_calls == []
    # Skipped by the zip check itself, and said so exactly once.
    assert not documents.zip_container_ok(bomb, fmt)
    doc_log.clear()
    _parse(fmt, bomb)
    lines = [r.getMessage() for r in doc_log if str(bomb) in r.getMessage()]
    assert len(lines) == 1 and "inflates past" in lines[0], lines


@pytest.mark.parametrize(("fmt", "member"), _ZIP_FORMATS)
def test_high_ratio_member_is_skipped(tmp_path, library_calls, fmt, member):
    # 2 MB of zeros: far under the 64 MB total, but about 1000:1.
    f = tmp_path / f"ratio.{fmt}"
    _write_zip(f, {"[Content_Types].xml": b"<Types/>", member: 2 << 20})
    assert _parse(fmt, f).sections == []
    assert library_calls == []


@pytest.mark.parametrize(("fmt", "member"), _ZIP_FORMATS)
def test_too_many_members_is_skipped(tmp_path, library_calls, fmt, member):
    f = tmp_path / f"many.{fmt}"
    _write_zip(f, {f"x/{i}.xml": b"<a/>" for i in range(2_001)})
    assert _parse(fmt, f).sections == []
    assert library_calls == []


@pytest.mark.parametrize(("fmt", "member"), _ZIP_FORMATS)
def test_total_uncompressed_cap_without_a_ratio_signature(
    tmp_path, library_calls, monkeypatch, fmt, member
):
    import os

    # Incompressible members: no ratio alarm, only the total size cap.
    monkeypatch.setattr(documents, "_MAX_ZIP_UNCOMPRESSED", 2 << 20)
    f = tmp_path / f"big.{fmt}"
    _write_zip(f, {f"m/{i}.bin": os.urandom(1 << 20) for i in range(3)})
    assert _parse(fmt, f).sections == []
    assert library_calls == []


def test_real_documents_pass_the_zip_check(tmp_path):
    for name, make in [
        ("c.docx", _make_docx),
        ("b.xlsx", _make_xlsx),
        ("d.pptx", _make_pptx),
    ]:
        make(tmp_path / name)
        assert documents.zip_container_ok(tmp_path / name, name[-4:])


def test_hard_input_cap_applies_even_when_forced(tmp_path, monkeypatch):
    f = tmp_path / "doc.pdf"
    f.write_bytes(_mini_pdf())
    monkeypatch.setattr(documents, "_MAX_INPUT_BYTES", 100)
    assert PdfParser().parse(f).sections == []


# ---------------------------------------------------------------------------
# a library missing from the install
# ---------------------------------------------------------------------------


def test_missing_library_skips_with_one_log_line(tmp_path, monkeypatch, doc_log):
    f = tmp_path / "deck.pptx"
    _make_pptx(f)
    monkeypatch.setitem(sys.modules, "pptx", None)  # import pptx raises
    monkeypatch.setattr(documents, "_missing_logged", set())
    first = PptxParser().parse(f)
    second = PptxParser().parse(f)
    assert first.sections == [] and second.sections == []
    assert documents._missing_logged == {"pptx"}
    lines = [r for r in doc_log if "pptx files skipped" in r.getMessage()]
    assert len(lines) == 1


# ---------------------------------------------------------------------------
# end to end: index, search, and the incremental pick-up
# ---------------------------------------------------------------------------


def _git(root, *args):
    subprocess.run(
        ["git", "-c", "user.name=t", "-c", "user.email=t@t", *args],
        cwd=root,
        check=True,
        capture_output=True,
    )


def _write_all(root):
    (root / "a.pdf").write_bytes(_mini_pdf("Pelican invoice"))
    _make_docx(root / "contract.docx")
    _make_xlsx(root / "book.xlsx")
    _make_pptx(root / "deck.pptx")
    (root / "t.csv").write_text("aardvark_id,score\n1,2\n", encoding="utf-8")


def _search(root, term):
    from codegraph.core.fts import fts_search, get_fts_conn

    return [
        h.file_path
        for h in fts_search(get_fts_conn(root), term, limit=20)
        if h.kind == "md_section"
    ]


@pytest.fixture
def fresh_db():
    from codegraph.core.db import reset_connection

    reset_connection()
    yield
    reset_connection()


def test_indexed_documents_are_searchable(tmp_path, fresh_db):
    from codegraph.indexer import index_repo

    _git(tmp_path, "init", "-q")
    _write_all(tmp_path)
    stats = index_repo(str(tmp_path))
    assert stats["errors"] == 0

    for term, name in [
        ("pelican", "a.pdf"),
        ("payment", "contract.docx"),
        ("donor_email", "book.xlsx"),
        ("walrus", "deck.pptx"),
        ("platypus", "deck.pptx"),
        ("aardvark_id", "t.csv"),
    ]:
        assert any(p.endswith(name) for p in _search(tmp_path, term)), term


def test_incremental_index_picks_up_previously_skipped_documents(
    tmp_path, fresh_db, monkeypatch
):
    """An index written by a cgh without these parsers never stored the
    document files. The incremental pass re-indexes every supported file in
    HEAD that the graph does not hold, so they come in without a graph
    format bump (which would force a full re-parse of every repo)."""
    from codegraph.indexer import incremental_reindex, index_repo

    _git(tmp_path, "init", "-q")
    _write_all(tmp_path)
    (tmp_path / "m.py").write_text("def f():\n    return 1\n", encoding="utf-8")
    _git(tmp_path, "add", "-A")
    _git(tmp_path, "commit", "-qm", "init")

    # Simulate the older core: no document parsers registered.
    doc_exts = (".pdf", ".docx", ".xlsx", ".xlsm", ".pptx")
    saved = {e: parsers._REGISTRY.pop(e) for e in doc_exts}
    for e in doc_exts:
        parsers._INSTANCES.pop(e, None)
    try:
        index_repo(str(tmp_path))
        assert _search(tmp_path, "walrus") == []
    finally:
        parsers._REGISTRY.update(saved)

    result = incremental_reindex(str(tmp_path))
    assert result.get("mode") != "fallback_full"
    assert any(p.endswith("deck.pptx") for p in _search(tmp_path, "walrus"))
    assert any(p.endswith("contract.docx") for p in _search(tmp_path, "payment"))
    assert any(p.endswith("a.pdf") for p in _search(tmp_path, "pelican"))
